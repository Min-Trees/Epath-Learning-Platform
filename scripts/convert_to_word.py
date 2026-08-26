# -*- coding: utf-8 -*-
# Script chuyen doi Markdown sang Word
# Yeu cau: pip install python-docx

import sys
import io

# Set UTF-8 output
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def convert_markdown_to_word(md_file, docx_file):
    """Chuyen doi file Markdown sang Word"""

    doc = Document()

    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    for i, line in enumerate(lines):
        if not line.strip():
            continue

        # Tieu de H1
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)

        # Tieu de H2
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)

        # Tieu de H3
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)

        # Tieu de H4
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)

        # Horizontal rule
        elif line.startswith('---') or line.startswith('***'):
            doc.add_paragraph('_' * 60)

        # Danh sach khong thu tu
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(text)

        # Danh sach co thu tu
        elif re.match(r'^\d+\.\s', line):
            match = re.match(r'^(\d+)\.\s(.*)', line)
            if match:
                text = match.group(2)
                p = doc.add_paragraph(style='List Number')
                p.add_run(text)

        # Bang
        elif line.startswith('|'):
            table_lines = [line]
            j = i + 1
            while j < len(lines) and lines[j].startswith('|'):
                table_lines.append(lines[j])
                j += 1

            # Loc bo dong phan cach
            data_rows = [l for l in table_lines if '---' not in l]

            if len(data_rows) >= 2:
                cols = len([c for c in data_rows[0].split('|') if c.strip()])

                if cols > 0:
                    table = doc.add_table(rows=len(data_rows), cols=cols)
                    table.style = 'Table Grid'

                    for row_idx, row_data in enumerate(data_rows):
                        cells = [c.strip() for c in row_data.split('|') if c.strip()]
                        for col_idx, cell_text in enumerate(cells[:cols]):
                            if col_idx < cols:
                                cell = table.rows[row_idx].cells[col_idx]
                                cell.text = cell_text

        # Code block
        elif line.startswith('```'):
            code_lines = []
            j = i + 1
            while j < len(lines) and not lines[j].startswith('```'):
                code_lines.append(lines[j])
                j += 1

            if code_lines:
                p = doc.add_paragraph()
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                p.paragraph_format.left_indent = Inches(0.3)

        # Text binh thuong
        else:
            text = line
            # Bo marker hinh anh [HINH X: ...]
            text = re.sub(r'\[HINH \d+:.*?\]', '', text)
            text = re.sub(r'\[.*?\]\(.*?\)', '', text)

            if text.strip():
                doc.add_paragraph(text)

    doc.save(docx_file)
    print("Da chuyen doi thanh cong!")

if __name__ == "__main__":
    md_file = "Noi_Dung_Slide_LP_Training_Hub.md"
    docx_file = "Huong_Dan_LP_Training_Hub.docx"

    try:
        convert_markdown_to_word(md_file, docx_file)
        print("Hoan thanh!")
    except Exception as e:
        print("Loi: " + str(e))
