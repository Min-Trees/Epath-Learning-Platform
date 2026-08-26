"""
Script tao PowerPoint huong dan he thong LP Training Hub
Thiet ke chuan doanh nghiep - Co cho chen anh
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Tao presentation - widescreen 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Mau sac thuong hieu - Corporate Blue
BRAND_BLUE = RGBColor(0, 82, 147)       # #005293 - Xanh dam
ACCENT_CYAN = RGBColor(0, 153, 204)    # #0099CC - Xanh cyan accent  
ACCENT_ORANGE = RGBColor(255, 133, 0)   # #FF8500 - Cam accent
LIGHT_BLUE = RGBColor(230, 244, 255)    # #E6F4FF - Xanh nhat
LIGHT_GRAY = RGBColor(248, 248, 248)    # #F8F8F8
DARK_GRAY = RGBColor(51, 51, 51)       # #333333
MID_GRAY = RGBColor(102, 102, 102)     # #666666
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
SUCCESS_GREEN = RGBColor(34, 139, 34)  # #228B22
WARNING_ORANGE = RGBColor(255, 165, 0) # #FFA500


def add_cover_slide(prs, title, subtitle, department="Phong Nhan su - HR"):
    """Slide bia - Cover page chuyen nghiep"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background gradient effect
    # Top accent bar
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = BRAND_BLUE
    top_bar.line.fill.background()
    
    # Left vertical bar
    left_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = BRAND_BLUE
    left_bar.line.fill.background()
    
    # Accent stripe
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), 0, Inches(0.08), prs.slide_height
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_CYAN
    accent.line.fill.background()
    
    # Logo placeholder
    logo_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.3), Inches(0.4), Inches(2.7), Inches(1.1)
    )
    logo_box.fill.solid()
    logo_box.fill.fore_color.rgb = RGBColor(240, 248, 255)
    logo_box.line.color.rgb = BRAND_BLUE
    logo_box.line.width = Pt(1.5)
    
    logo_text = slide.shapes.add_textbox(Inches(10.4), Inches(0.65), Inches(2.5), Inches(0.6))
    tf = logo_text.text_frame
    p = tf.paragraphs[0]
    p.text = "[COMPANY LOGO]"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    # Main title
    title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(10.5), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    p.alignment = PP_ALIGN.LEFT
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.7), Inches(10), Inches(0.9))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(26)
    p.font.color.rgb = DARK_GRAY
    p.alignment = PP_ALIGN.LEFT
    
    # Decorative line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(4.8), Inches(4), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_ORANGE
    line.line.fill.background()
    
    # Department
    dept_box = slide.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(10), Inches(0.5))
    tf = dept_box.text_frame
    p = tf.paragraphs[0]
    p.text = department
    p.font.size = Pt(16)
    p.font.color.rgb = MID_GRAY
    p.alignment = PP_ALIGN.LEFT
    
    # Version info
    ver_box = slide.shapes.add_textbox(Inches(1.2), Inches(5.7), Inches(10), Inches(0.4))
    tf = ver_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Version 1.0 - 2024"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(180, 180, 180)
    p.alignment = PP_ALIGN.LEFT
    
    # Bottom bar
    bottom_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.15), prs.slide_width, Inches(0.35)
    )
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = BRAND_BLUE
    bottom_bar.line.fill.background()
    
    return slide


def add_content_slide(prs, title, content_list, image_placeholder=True, section_num=""):
    """Slide noi dung co cho chen anh"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header bar
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.25)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = BRAND_BLUE
    header.line.fill.background()
    
    # Left accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.25), Inches(0.12), Inches(6.25)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_ORANGE
    accent.line.fill.background()
    
    # Section number badge
    if section_num:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.4), Inches(0.32), Inches(0.6), Inches(0.6)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_CYAN
        badge.line.fill.background()
        
        badge_text = slide.shapes.add_textbox(Inches(0.4), Inches(0.38), Inches(0.6), Inches(0.5))
        tf = badge_text.text_frame
        p = tf.paragraphs[0]
        p.text = section_num
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.1), Inches(0.3), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Content area
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(7.8), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        if isinstance(item, dict):
            text = item.get('text', '')
            prefix = "▶ " if item.get('bold') else "• "
            p.text = prefix + text
            p.font.size = Pt(item.get('size', 18))
            p.font.bold = item.get('bold', False)
            p.font.color.rgb = BRAND_BLUE if item.get('bold') else BLACK
            p.space_after = Pt(item.get('space', 10))
        elif isinstance(item, str):
            if item == "":
                p.text = ""
                p.space_after = Pt(6)
            else:
                p.text = "• " + item
                p.font.size = Pt(18)
                p.font.color.rgb = BLACK
                p.space_after = Pt(10)
        else:
            p.text = "• " + str(item)
            p.font.size = Pt(18)
            p.font.color.rgb = BLACK
            p.space_after = Pt(10)
    
    # Image placeholder
    if image_placeholder:
        img_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.5), Inches(4.5), Inches(5.3)
        )
        img_box.fill.solid()
        img_box.fill.fore_color.rgb = LIGHT_GRAY
        img_box.line.color.rgb = RGBColor(200, 210, 220)
        img_box.line.width = Pt(2)
        
        # Icon/image indicator
        icon_box = slide.shapes.add_textbox(Inches(8.5), Inches(3.6), Inches(4.5), Inches(0.8))
        tf = icon_box.text_frame
        p = tf.paragraphs[0]
        p.text = "[HINH ANH MO TA]"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(160, 170, 180)
        p.alignment = PP_ALIGN.CENTER
        
        img_desc = slide.shapes.add_textbox(Inches(8.5), Inches(4.3), Inches(4.5), Inches(1))
        tf = img_desc.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Chen hinh minh hoa\ntai day"
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(180, 180, 180)
        p.alignment = PP_ALIGN.CENTER
    
    # Footer
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.15), prs.slide_width, Inches(0.35)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = BRAND_BLUE
    footer.line.fill.background()
    
    return slide


def add_step_slide(prs, title, steps, section_num=""):
    """Slide huong dan theo buoc"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.25)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = BRAND_BLUE
    header.line.fill.background()
    
    # Left accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.25), Inches(0.12), Inches(6.25)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_ORANGE
    accent.line.fill.background()
    
    # Section badge
    if section_num:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.4), Inches(0.32), Inches(0.6), Inches(0.6)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_CYAN
        badge.line.fill.background()
        
        badge_text = slide.shapes.add_textbox(Inches(0.4), Inches(0.38), Inches(0.6), Inches(0.5))
        tf = badge_text.text_frame
        p = tf.paragraphs[0]
        p.text = section_num
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.1), Inches(0.3), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Steps
    num_steps = len(steps)
    step_height = min(Inches(0.95), Inches(5.2 / num_steps))
    start_y = Inches(1.5)
    
    for i, step in enumerate(steps):
        y = start_y + (i * step_height)
        
        # Step number
        num_box = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.5), y, Inches(0.55), Inches(0.55)
        )
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = BRAND_BLUE
        num_box.line.fill.background()
        
        num_text = slide.shapes.add_textbox(Inches(0.5), y + Inches(0.08), Inches(0.55), Inches(0.4))
        tf = num_text.text_frame
        p = tf.paragraphs[0]
        p.text = str(i + 1)
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        # Step content
        step_box = slide.shapes.add_textbox(Inches(1.2), y + Inches(0.05), Inches(6.8), Inches(0.8))
        tf = step_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(17)
        p.font.color.rgb = BLACK
        
        # Connector line
        if i < num_steps - 1:
            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0.76), y + Inches(0.58), Inches(0.03), Inches(0.35)
            )
            line.fill.solid()
            line.fill.fore_color.rgb = RGBColor(200, 210, 220)
            line.line.fill.background()
    
    # Image placeholder
    img_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.5), Inches(4.5), Inches(5.3)
    )
    img_box.fill.solid()
    img_box.fill.fore_color.rgb = LIGHT_GRAY
    img_box.line.color.rgb = RGBColor(200, 210, 220)
    img_box.line.width = Pt(2)
    
    icon_box = slide.shapes.add_textbox(Inches(8.5), Inches(3.6), Inches(4.5), Inches(0.8))
    tf = icon_box.text_frame
    p = tf.paragraphs[0]
    p.text = "[HINH ANH MO TA]"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(160, 170, 180)
    p.alignment = PP_ALIGN.CENTER
    
    img_desc = slide.shapes.add_textbox(Inches(8.5), Inches(4.3), Inches(4.5), Inches(1))
    tf = img_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Chen hinh minh hoa\ntai day"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(180, 180, 180)
    p.alignment = PP_ALIGN.CENTER
    
    # Footer
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.15), prs.slide_width, Inches(0.35)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = BRAND_BLUE
    footer.line.fill.background()
    
    return slide


def add_table_slide(prs, title, headers, rows, section_num=""):
    """Slide voi bang thong tin"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.25)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = BRAND_BLUE
    header.line.fill.background()
    
    # Left accent
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.25), Inches(0.12), Inches(6.25)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_ORANGE
    accent.line.fill.background()
    
    # Section badge
    if section_num:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.4), Inches(0.32), Inches(0.6), Inches(0.6)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_CYAN
        badge.line.fill.background()
        
        badge_text = slide.shapes.add_textbox(Inches(0.4), Inches(0.38), Inches(0.6), Inches(0.5))
        tf = badge_text.text_frame
        p = tf.paragraphs[0]
        p.text = section_num
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.1), Inches(0.3), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Table
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_width = Inches(7.5)
    table_height = Inches(min(num_rows * 0.55, 5.2))
    
    table = slide.shapes.add_table(
        num_rows, num_cols, Inches(0.5), Inches(1.55), table_width, table_height
    ).table
    
    # Header row
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND_BLUE
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.size = Pt(14)
        p.alignment = PP_ALIGN.CENTER
    
    # Data rows
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_text)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(13)
            p.font.color.rgb = BLACK
            if row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(245, 248, 252)
    
    # Footer note
    note_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(7.5), Inches(0.4))
    tf = note_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Nhan vao dong de xem chi tiet"
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = MID_GRAY
    
    # Footer
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.15), prs.slide_width, Inches(0.35)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = BRAND_BLUE
    footer.line.fill.background()
    
    return slide


def add_end_slide(prs, title, message, contact=""):
    """Slide ket thuc"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BRAND_BLUE
    bg.line.fill.background()
    
    # Decorative circles
    circle1 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(-1.5), Inches(-1.5), Inches(5), Inches(5)
    )
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = RGBColor(0, 102, 180)
    circle1.line.fill.background()
    
    circle2 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(10.5), Inches(4.5), Inches(4), Inches(4)
    )
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = RGBColor(0, 102, 180)
    circle2.line.fill.background()
    
    # Main text
    main_box = slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11), Inches(1.3))
    tf = main_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Message
    msg_box = slide.shapes.add_textbox(Inches(1), Inches(3.9), Inches(11), Inches(1))
    tf = msg_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = message
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(200, 220, 245)
    p.alignment = PP_ALIGN.CENTER
    
    # Decorative line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(5), Inches(2.3), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_ORANGE
    line.line.fill.background()
    
    # Contact
    if contact:
        contact_box = slide.shapes.add_textbox(Inches(1), Inches(5.4), Inches(11), Inches(0.8))
        tf = contact_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = contact
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(180, 200, 230)
        p.alignment = PP_ALIGN.CENTER
    
    return slide


# ============ TAO CAC SLIDE ============

# SLIDE 1: Bia
add_cover_slide(
    prs,
    "LP TRAINING HUB",
    "Huong dan su dung he thong dao tao truc tuyen",
    "Phong Nhan su - HR Department"
)

# SLIDE 2: Muc luc
add_content_slide(prs, "Muc luc", [
    "01. Gioi thieu he thong",
    "02. Dang nhap he thong",
    "03. Giao dien trang chu",
    "04. Chuong trinh cua toi",
    "05. Noi dung bai hoc - Video",
    "06. Tai lieu dinh kem",
    "07. Bai kiem tra",
    "08. Tien do hoc tap",
    "09. Khoa hoc cong khai",
    "10. Gui yeu cau ho tro",
    "11. Xu ly su co thuong gap",
    "12. Thong tin lien he",
])

# SLIDE 3: Gioi thieu
add_content_slide(prs, "01. Gioi thieu LP Training Hub", [
    {"text": "LP Training Hub la gi?", "bold": True, "size": 18},
    "He thong dao tao truc tuyen noi bo",
    "Giup nhan vien hoc tap moi luc, moi noi",
    "Theo doi tien do hoc tap",
    "Nhan chung chi khi hoan thanh",
    "",
    {"text": "Cac chuc nang chinh", "bold": True, "size": 18},
    "Trang chu - Tong quan tien do",
    "Chuong trinh cua toi - Khoa hoc duoc gan",
    "Khoa hoc - Khoa hoc cong khai",
    "Ho so - Thong tin ca nhan",
    "Ho tro - Gui yeu cau",
], section_num="01")

# SLIDE 4: Dang nhap
add_step_slide(prs, "02. Dang nhap He thong", [
    "Mo trinh duyet web (Chrome, Edge, Firefox...)",
    "Truy cap dia chi: [URL he thong LP Training Hub]",
    "Nhap email cong ty va mat khau",
    "Nhan nut 'Dang nhap'",
    "Xac thuc 2 buoc neu duoc yeu cau",
], section_num="02")

# SLIDE 5: Giao dien trang chu
add_content_slide(prs, "03. Giao dien Trang chu - Dashboard", [
    {"text": "Thong tin hien thi", "bold": True, "size": 18},
    "Xin chao + Ten nhan vien",
    "So lieu thong ke: Chuong trinh duoc gan, Dang hoc, Da hoan thanh",
    "Tien do trung binh (%)",
    "",
    {"text": "Cac thanh phan", "bold": True, "size": 18},
    "Menu dieu huong ben trai",
    "Noi dung chinh o giua",
    "Thong bao o goc phai tren",
    "",
    {"text": "Thao tac nhanh", "bold": True, "size": 18},
    "Xem tat ca chuong trinh",
    "Xem tien do cua toi",
], section_num="03")

# SLIDE 6: Chuong trinh cua toi
add_content_slide(prs, "04. Chuong trinh cua toi", [
    {"text": "Xem danh sach", "bold": True, "size": 18},
    "Cac khoa hoc duoc admin gan",
    "Phan loai: Dang hoc, Da hoan thanh",
    "Hien thi tien do cua tung chuong trinh",
    "",
    {"text": "Thong tin hien thi", "bold": True, "size": 18},
    "Ten chuong trinh va mo ta",
    "Trang thai: Chua bat dau / Dang hoc / Hoan thanh",
    "Phan tram tien do",
    "So bai da hoc / Tong so bai",
    "",
    {"text": "Bat dau hoc", "bold": True, "size": 18},
    "Nhan 'Bat dau hoc' hoac 'Tiep tuc'",
    "He thong chuyen den trang bai hoc",
], section_num="04")

# SLIDE 7: Video bai hoc
add_step_slide(prs, "05. Noi dung BAI HOC - Video", [
    "Chon bai hoc tu danh sach trong chuong trinh",
    "Video bat dau phat o giua man hinh",
    "Su dung cac nut dieu khien: Play/Pause, Am luong",
    "Tua video bang thanh tien trinh",
    "Chon toc do phat: 0.5x, 1x, 1.25x, 1.5x, 2x",
    "Nhan 'Hoan thanh' sau khi xem xong",
], section_num="05")

# SLIDE 8: Tai lieu
add_content_slide(prs, "06. Tai lieu Dinh kem", [
    {"text": "Vi tri", "bold": True, "size": 18},
    "Nam o phia duoi video",
    "Hien thi khi bai hoc co tai lieu di kem",
    "",
    {"text": "Cac loai file ho tro", "bold": True, "size": 18},
    "PDF (.pdf) - Tai lieu dinh dang PDF",
    "Word (.doc, .docx) - Van ban",
    "Excel (.xls, .xlsx) - Bang tinh",
    "PowerPoint (.ppt, .pptx) - Trinh chieu",
    "",
    {"text": "Cach tai", "bold": True, "size": 18},
    "Nhan vao ten file hoac nut tai ve",
    "Mot so file co the xem truc tiep",
], section_num="06")

# SLIDE 9: Bai kiem tra
add_content_slide(prs, "07. Bai Kiem tra", [
    {"text": "Khi nao lam bai thi", "bold": True, "size": 18},
    "Sau khi hoan thanh cac bai hoc",
    "Hoac tai bai hoc cuoi co bai kiem tra",
    "",
    {"text": "Cac buoc lam bai", "bold": True, "size": 18},
    "Doc ky de bai va huong dan",
    "Tra loi tung cau hoi",
    "Co the danh dau de xem lai",
    "Nop bai khi hoan thanh",
    "",
    {"text": "Xem ket qua", "bold": True, "size": 18},
    "Diem so hien thi ngay sau khi nop",
    "Xem dap an dung/sai",
    "Co the lam lai neu cho phep",
], section_num="07")

# SLIDE 10: Tien do hoc tap
add_content_slide(prs, "08. Tien do Hoc tap", [
    {"text": "Xem tien do ca nhan", "bold": True, "size": 18},
    "Vao 'Tien do cua toi' tu menu",
    "Tien do trung binh tat ca khoa hoc",
    "Danh sach khoa hoc theo trang thai",
    "",
    {"text": "Bieu do thong ke", "bold": True, "size": 18},
    "Thanh phan tram tien do tong quan",
    "Danh sach chi tiet tung khoa",
    "",
    {"text": "Chung chi", "bold": True, "size": 18},
    "Nhan chung chi khi hoan thanh 100%",
    "Co the tai ve PDF hoac chia se",
], section_num="08")

# SLIDE 11: Khoa hoc cong khai
add_content_slide(prs, "09. Khoa hoc Cong khai", [
    {"text": "Khac biet voi Chuong trinh", "bold": True, "size": 18},
    "Tu dang ky (khong phai admin gan)",
    "Tuy chon (khong bat buoc)",
    "",
    {"text": "Tim kiem va loc", "bold": True, "size": 18},
    "Tim kiem theo tu khoa",
    "Loc theo trinh do: Co ban, Trung binh, Nang cao",
    "Loc theo danh muc: CNTT, HR, Sales, Marketing...",
    "Xep sep: Moi nhat, Pho bien, Danh gia cao",
    "",
    {"text": "Dang ky", "bold": True, "size": 18},
    "Nhan 'Dang ky' hoac 'Tham gia'",
    "Khoa hoc duoc them vao danh sach",
], section_num="09")

# SLIDE 12: Gui yeu cau ho tro
add_step_slide(prs, "10. Gui Yeu cau Ho tro", [
    "Vao menu 'Ho tro'",
    "Nhan 'Tao yeu cau moi'",
    "Dien thong tin: Tieu de, Danh muc, Muc do uu tien",
    "Viet mo ta chi tiet van de gap phai",
    "Dinh kem anh chup man hinh (neu can)",
    "Nhan 'Gui yeu cau'",
], section_num="10")

# SLIDE 13: Trang thai yeu cau
add_table_slide(prs, "Trang thai Yeu cau Ho tro", 
    ["Trang thai", "Nghia", "Han dong tiep theo"],
    [
        ["Moi", "Vua gui, cho xu ly", "Cho nhan vien ho tro xem xet"],
        ["Dang xu ly", "Duoc nhan, dang xem xet", "Cho phan hoi tu ho tro"],
        ["Da tra loi", "Co phan hoi", "Kiem tra va phan hoi lai neu can"],
        ["Da giai quyet", "Van de da xu ly xong", "Xac nhan dong yeu cau"],
        ["Da dong", "Yeu cau hoan tat", "Da ket thuc"],
    ], section_num="10")

# SLIDE 14: Xu ly su co
add_content_slide(prs, "11. Xu ly Su co Thuong gap", [
    {"text": "Video khong phat", "bold": True, "size": 18},
    "Kiem tra ket noi internet",
    "Refresh trang hoac xoa cache trinh duyet",
    "Thu trinh duyet khac",
    "",
    {"text": "Khong dang nhap duoc", "bold": True, "size": 18},
    "Kiem tra email va mat khau",
    "Nhan 'Quen mat khau' neu can",
    "Lien he IT neu tai khoan bi khoa",
    "",
    {"text": "Bai thi mat ket qua", "bold": True, "size": 18},
    "Khong tat tab khi dang lam bai",
    "Thu lam lai neu con luot",
], section_num="11")

# SLIDE 15: Thong tin lien he
add_content_slide(prs, "12. Thong tin Lien he Ho tro", [
    {"text": "Bo phan ho tro", "bold": True, "size": 18},
    "Phong Nhan su (HR): Khoa hoc, Chung chi",
    "Phong CNTT (IT): Loi ky thuat, He thong",
    "",
    {"text": "Cach lien he", "bold": True, "size": 18},
    "Qua he thong: Menu Ho tro > Tao yeu cau",
    "Qua email: hr@company.com",
    "Truc tiep: Den phong ban lien quan",
    "",
    {"text": "Thoi gian ho tro", "bold": True, "size": 18},
    "Gio hanh chinh: 8:00 - 17:30 (T2 - T6)",
    "Khong lam viec: Chu nhat, ngay le",
], section_num="12")

# SLIDE 16: Ket thuc
add_end_slide(
    prs,
    "CAM ON",
    "Chuc cac ban hoc tap hieu qua!",
    "Lien he: hr@company.com | it-support@company.com"
)

# Luu file
output_path = "Slide_LP_Training_Hub_Huong_Dan.pptx"
prs.save(output_path)
print(f"Da tao file: {output_path}")
print(f"Tong so slide: {len(prs.slides)}")
